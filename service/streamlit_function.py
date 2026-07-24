from service.function import parse_song_form, render_song
from pptx import Presentation
from bs4 import BeautifulSoup
from urllib.parse import quote

import matplotlib.pyplot as plt
import pandas as pd
import numpy as np
import re
import streamlit as st
import requests


def load_css(path: str):
    with open(path) as f:
        st.markdown(f"<style>{f.read()}</style>", unsafe_allow_html=True)


def reset_session():
    st.session_state.part_count = 3
    st.session_state.reset_counter += 1

    st.session_state.pop("extracted_text", None)

    st.toast("모든 입력이 초기화되었습니다", icon="🔄")
    st.rerun()


def collect_parts_from_session() -> dict[str, str]:
    """
    session_state에 저장된 파트명/가사를 읽어서
    { "A": "가사", "B": "가사" } 형태로 반환
    """
    parts: dict[str, str] = {}

    reset_counter = st.session_state.get("reset_counter", 0)
    part_count = st.session_state.get("part_count", 0)

    for i in range(part_count):
        name = st.session_state.get(f"part_name_{i}_{reset_counter}", "").strip()

        lyrics = st.session_state.get(f"part_lyrics_{i}_{reset_counter}", "").strip()

        if name and lyrics:
            parts[name] = lyrics

    return parts


def export_holiday(song_form: str) -> str:
    parts = collect_parts_from_session()

    parsed = parse_song_form(song_form)
    output_lines = []

    for token in parsed:
        if token.startswith("(") and token.endswith(")"):
            continue

        if token in parts:
            output_lines.append(parts[token].strip())
            output_lines.append("")

    result = "\n".join(output_lines)
    result = re.sub(r"//+", "", result).strip()

    formatted = f"""{st.session_state.get("song_title", "")}

{result}
"""

    # pyperclip.copy(formatted)

    st.session_state.extracted_text = formatted
    st.toast("생성 완료 ✅", icon="📋")


def export_retreat(song_form: str):
    parts_dict = collect_parts_from_session()

    formatted = f"""{{
    "title": "{st.session_state.get("song_title", "")}",
    "parts": {{
"""

    for k, v in parts_dict.items():
        indented_v = indent_text(v, spaces=18)

        formatted += f'''        "{k}": """
{indented_v}
""",
'''

    formatted += f"""    }},
    "song_form": "{song_form}",
}},
"""

    # pyperclip.copy(formatted)
    st.session_state.extracted_text = formatted
    st.toast("생성 완료 ✅", icon="📋")


def indent_text(text: str, spaces: int = 8) -> str:
    indent = " " * spaces
    return "\n".join(
        indent + line if line.strip() else line for line in text.splitlines()
    )


def _remove_all_slides(prs):
    slide_id_list = prs.slides._sldIdLst

    for slide_id in list(slide_id_list):
        prs.part.drop_rel(slide_id.rId)
        slide_id_list.remove(slide_id)


def ppt_save(song_list, path, template=None):
    prs = Presentation(template) if template else Presentation()

    if template:
        _remove_all_slides(prs)

    for song in song_list:
        render_song(prs, song, apply_default_background=not template)

    prs.save(path)


def crawl_lyrics(song_name: str, limit: int = 8):
    url = f"https://music.bugs.co.kr/search/lyrics?q={quote(song_name)}"
    headers = {"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7)"}

    res = requests.get(url, headers=headers, timeout=10)
    res.raise_for_status()
    soup = BeautifulSoup(res.text, "html.parser")

    results = []

    rows = soup.select('tr[rowtype="lyrics"]')

    print(rows)

    for row in rows:
        if len(results) >= limit:
            break

        # track_id (진짜 곡 ID)
        track_id = row.get("trackid", "").strip()

        # 제목
        title_a = row.select_one("a[title]")
        title = title_a.get_text(" ", strip=True) if title_a else ""

        # 아티스트
        artist_a = row.select_one('a[href*="/artist/"]')
        artist = artist_a.get_text(" ", strip=True) if artist_a else ""

        if track_id and title:
            results.append(
                {
                    "track_id": track_id,
                    "title": title,
                    "artist": artist,
                }
            )
    return results


def crawl_track_lyrics(track_id: str) -> str:
    url = f"https://music.bugs.co.kr/track/{track_id}"
    headers = {"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7)"}

    res = requests.get(url, headers=headers, timeout=10)
    res.raise_for_status()

    soup = BeautifulSoup(res.text, "html.parser")

    # div.lyricsContainer 안에 xmp에 가사가 들어있는 구조
    xmp = soup.select_one("div.lyricsContainer xmp")
    if not xmp:
        return ""

    lyrics = xmp.get_text("\n", strip=True)
    return lyrics


def sync_lyrics_with_track():
    """
    track_id 변경을 감지해서
    해당 곡의 가사를 세션에 동기화한다.
    """

    current_track_id = st.session_state.get("track_id")
    prev_track_id = st.session_state.get("prev_track_id")

    if current_track_id and current_track_id != prev_track_id:
        try:
            st.session_state.lyrics_text = crawl_track_lyrics(current_track_id)
            st.session_state.prev_track_id = current_track_id
        except Exception as e:
            st.toast(f"가사 불러오기 실패: {e}", icon="❌")


def listup_lyrics_result(results):
    if not results:
        st.info("검색 결과가 여기에 표시됩니다. (상위 8개만 표시)")
    else:
        for idx, r in enumerate(results):
            st.markdown(
                f"""
                <div class="song-card-wrapper">
                    <div class="song-card">
                        <div class="song-title">{r["title"]}</div>
                        <div class="song-artist">{r["artist"]}</div>
                    </div>
                </div>
                """,
                unsafe_allow_html=True,
            )

            if st.button(
                "select",
                key=f"search_select_{idx}",
                use_container_width=True,
            ):
                st.session_state.song_title = r["title"]
                st.session_state.song_artist = r["artist"]
                st.session_state.track_id = r["track_id"]
                st.rerun()


def plot_chroma_histogram(chroma_vec):
    notes = ["C", "C#", "D", "D#", "E", "F", "F#", "G", "G#", "A", "A#", "B"]

    fig, ax = plt.subplots(figsize=(8, 3))
    ax.bar(notes, chroma_vec)
    ax.set_title("Pitch Class Histogram (Chroma)")
    ax.set_ylabel("Normalized Energy")
    ax.set_xlabel("Note")

    return fig


def build_key_ranking_table(ranked_candidates, top_n=10):
    rows = []
    for i, (key, mode, score) in enumerate(ranked_candidates[:top_n], start=1):
        rows.append(
            {"Rank": i, "Key": key, "Mode": mode, "Correlation": round(score, 3)}
        )

    return pd.DataFrame(rows)


def collect_auto_maker_parts() -> dict[str, str]:
    parts: dict[str, str] = {}
    reset_counter = st.session_state.get("auto_maker_reset_counter", 0)
    part_count = st.session_state.get("auto_maker_part_count", 0)

    for i in range(part_count):
        name = st.session_state.get(
            f"auto_maker_part_name_{i}_{reset_counter}", ""
        ).strip()
        lyrics = st.session_state.get(
            f"auto_maker_part_lyrics_{i}_{reset_counter}", ""
        ).strip()
        if name and lyrics:
            parts[name] = lyrics
    return parts


def build_sunday_text(song_form: str, parts: dict[str, str], title: str = "") -> str:
    tokens = parse_song_form(song_form)
    output_lines: list[str] = []

    for token in tokens:
        if token.startswith("(") and token.endswith(")"):
            continue
        text = parts.get(token, "").strip()
        if text:
            output_lines.append(text)
            output_lines.append("")

    result = "\n".join(output_lines)
    result = re.sub(r"//+", "", result).strip()
    if title:
        return f"{title}\n\n{result}".strip()
    return result


def handle_generate_action(
    generate_clicked: bool,
    request_openai_autofill,
) -> None:
    if not generate_clicked:
        return

    selected_song = st.session_state.get("auto_maker_selected_song")
    full_lyrics = st.session_state.get("auto_maker_lyrics_text", "").strip()

    if not selected_song:
        st.toast("먼저 상단 곡 박스에서 곡을 선택해주세요.", icon="⚠️")
        return
    if not full_lyrics:
        st.toast("가사 내용이 비어 있습니다.", icon="⚠️")
        return

    with st.spinner("가사를 파트별로 자동 생성중..."):
        result = request_openai_autofill(selected_song, full_lyrics)

    st.session_state.auto_maker_autofill_result = result
    if "error" in result:
        st.toast(f"자동생성 실패: {result.get('error')}", icon="❌")
        return

    st.toast("자동생성 완료", icon="✅")
    st.session_state.auto_maker_autofill_apply = result
    st.rerun()
